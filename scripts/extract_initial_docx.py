"""Extract the styled OPIc DOCX into a private, review-first JSON import package."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document


def clean(value: str | None) -> str:
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]+", " ", value or "")).strip()


def prefixed(lines: list[str], prefix: str) -> str:
    for line in lines:
        if line.startswith(prefix):
            return clean(line[len(prefix) :])
    return ""


def parse_heading(text: str) -> tuple[str | None, str, str, list[str]]:
    match = re.match(r"^([A-Z]?Q?\d+)\s*[|·-]\s*([^|·]+?)(?:\s*[|·]\s*(.+))?$", text, re.I)
    if not match:
        return None, "Unclassified", text, ["Question heading could not be parsed reliably."]
    return clean(match.group(1)), clean(match.group(2)), clean(match.group(3)) or clean(match.group(2)), []


def split_expressions(value: str) -> list[dict[str, str | None]]:
    result = []
    for part in re.split(r"\s*/\s*", value):
        if not clean(part):
            continue
        phrase, separator, meaning = part.partition("=")
        result.append({"phrase": clean(phrase), "meaning": clean(meaning) if separator else None})
    return result


def extract(source: Path) -> dict:
    document = Document(source)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    questions: list[dict] = []
    guides: list[dict] = []
    section = ""
    topic = ""
    heading = ""
    blocks: list[tuple[str, str]] = []
    guide_title = ""
    guide_lines: list[str] = []

    def flush_guide() -> None:
        nonlocal guide_lines
        body = clean("\n\n".join(guide_lines))
        if guide_title and body:
            guides.append(
                {
                    "clientId": f"guide-{len(guides) + 1}",
                    "selected": True,
                    "title": guide_title,
                    "section": section or None,
                    "body": body,
                    "sourceReference": f"Section: {' / '.join(filter(None, [section, guide_title]))}",
                    "reviewReasons": [],
                }
            )
        guide_lines = []

    def flush_question() -> None:
        nonlocal heading, blocks
        if not heading:
            return
        code, question_type, title, review = parse_heading(heading)
        prompts = [text for style, text in blocks if style == "Question"]
        answers = [text for style, text in blocks if style == "Answer"]
        analysis = [text for style, text in blocks if style == "Analysis"]
        small = [text for style, text in blocks if style == "Small"]
        translation = next((line for line in small if line.startswith("English translation")), "")
        source_reference = next((line for line in small if re.match(r"^(원자료 문항|추가·재구성 문항|Source)", line, re.I)), "")
        if not prompts:
            review.append("No Question-style paragraph was found.")
        if not answers:
            review.append("No answer was found; it will remain marked as missing.")
        if not source_reference:
            review.append("No source reference was detected.")
        if len(prompts) > 1:
            review.append("Multiple question paragraphs were preserved as possible subquestions.")
        explanation = [
            prefixed(analysis, "질문 의도"),
            prefixed(analysis, "구조 분석"),
            prefixed(analysis, "말하기 큐"),
            prefixed(analysis, "주의할 점"),
        ]
        questions.append(
            {
                "clientId": f"question-{len(questions) + 1}",
                "selected": True,
                "topic": topic or "Uncategorized",
                "questionCode": code,
                "title": title,
                "prompt": "\n\n".join(prompts),
                "promptTranslation": clean(re.sub(r"^English translation\s*", "", translation, flags=re.I)) or None,
                "questionType": question_type,
                "sourceReference": source_reference or None,
                "isSupplement": "추가·재구성" in source_reference or "supplement" in source_reference.lower(),
                "answers": [
                    {
                        "label": "Model answer",
                        "englishText": "\n\n".join(answers),
                        "koreanExplanation": "\n\n".join(filter(None, explanation)),
                    }
                ]
                if answers
                else [],
                "expressions": split_expressions(prefixed(analysis, "재사용 표현")),
                "notes": list(filter(None, [prefixed(analysis, "내용 근거"), prefixed(analysis, "MP")])),
                "reviewReasons": review,
            }
        )
        heading = ""
        blocks = []

    for paragraph in document.paragraphs:
        text = clean(paragraph.text)
        if not text:
            continue
        style = paragraph.style.name
        if style == "Heading 1":
            flush_question()
            flush_guide()
            section = text
            topic = ""
            guide_title = text
        elif style == "Heading 2":
            flush_question()
            flush_guide()
            topic = text
            guide_title = text
        elif style == "Heading 3":
            flush_question()
            flush_guide()
            heading = text
        elif heading:
            blocks.append((style, text))
        else:
            guide_lines.append(text)
    flush_question()
    flush_guide()

    for index, table in enumerate(document.tables, start=1):
        rows = [[clean(cell.text) for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        body = "\n".join(" | ".join(row) for row in rows)
        guides.append(
            {
                "clientId": f"guide-{len(guides) + 1}",
                "selected": True,
                "title": f"Source table {index}",
                "section": "Document tables",
                "body": body,
                "sourceReference": f"Table {index}",
                "reviewReasons": ["Table formatting was flattened to text and requires visual review."],
            }
        )

    normalized: dict[str, list[dict]] = {}
    for question in questions:
        key = re.sub(r"[^a-z0-9가-힣]+", " ", question["prompt"].lower()).strip()
        if key:
            normalized.setdefault(key, []).append(question)
    for duplicates in normalized.values():
        if len(duplicates) > 1:
            for question in duplicates:
                question["reviewReasons"].append("Possible duplicate prompt in this import.")

    extracted_at = datetime.now(timezone.utc).isoformat()
    review_count = sum(bool(item["reviewReasons"]) for item in questions + guides)
    return {
        "schemaVersion": 1,
        "source": {"filename": source.name, "sha256": source_hash, "extractedAt": extracted_at},
        "inventory": {
            "sourceFilename": source.name,
            "sha256": source_hash,
            "extractedAt": extracted_at,
            "paragraphCount": len(document.paragraphs),
            "tableCount": len(document.tables),
            "embeddedImageCount": len(document.inline_shapes),
            "questionCount": len(questions),
            "guideCount": len(guides),
            "selectedCount": len(questions) + len(guides),
            "reviewCount": review_count,
            "notConverted": [
                "Comments and tracked-change metadata are not imported.",
                "Floating or image-based text requires manual review.",
                "Table formatting is flattened in JSON; the original DOCX remains authoritative.",
                "PDF text and OCR are not extracted.",
            ],
            "messages": [],
        },
        "questions": questions,
        "guides": guides,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    package = extract(args.input.resolve())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(package, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(package["inventory"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
